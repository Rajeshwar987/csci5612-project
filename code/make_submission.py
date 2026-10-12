"""
make_submission.py  --  builds the Canvas submission document.

The assignment asks for a single URL placed in a .docx file, with the document
named for the student and the assignment.
"""

from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor

NAME = "Rajeshwar Rengeswaran"
COURSE = "CSCI 5612 - Machine Learning"
ASSIGNMENT = "Module 2 Project Part 2"
URL = "https://rajeshwar987.github.io/csci5612-project/"
REPO = "https://github.com/Rajeshwar987/csci5612-project"

OUT = Path.home() / "Downloads" / f"{NAME} Project A2.docx"

doc = Document()

h = doc.add_heading(f"{ASSIGNMENT} - Project Website", level=1)

p = doc.add_paragraph()
p.add_run(f"{NAME}\n").bold = True
p.add_run(f"{COURSE}\n")
p.add_run(ASSIGNMENT)

doc.add_paragraph()

def link_line(label, url):
    p = doc.add_paragraph()
    p.add_run(f"{label}\n").bold = True
    run = p.add_run(url)
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(0x0B, 0x51, 0x94)
    run.underline = True


link_line("Project website URL:", URL)
link_line("GitHub repository:", REPO)

doc.add_paragraph()

doc.add_paragraph(
    "Project topic: whether laboratory crash-test ratings predict the safety "
    "problems vehicle owners actually report once a vehicle is on the road."
)

doc.add_heading("What is new in Part 2", level=2)

for line in [
    "Introduction: extended to six paragraphs and a second image.",
    "Clustering tab: Overview covering partitional versus hierarchical methods "
    "and distance metrics, Data Prep explaining the unlabeled numeric "
    "requirement, Code, Results and Conclusions. k-means was swept from k=2 to "
    "k=15 and compared at k=3, 5 and 7; hierarchical clustering used cosine "
    "distance with average linkage. Both methods independently select k=7 "
    "(adjusted Rand index 0.638).",
    "PCA tab: Overview covering eigenvalues, eigenvectors and the problems of "
    "high-dimensional data, plus Data Prep, Code, Results and Conclusions. "
    "PC1 carries 24.2% of variance and nine components are needed for 90%.",
    "Figures added: silhouette curve and per-vehicle widths, a three-way k "
    "comparison, full and truncated dendrograms, a method agreement matrix, "
    "cluster profiles, scree and eigenvalue plots, a loadings matrix, a "
    "correlation circle, and a cluster projection.",
]:
    doc.add_paragraph(line, style="List Bullet")

doc.add_paragraph()
doc.add_paragraph(
    "All analysis code and all raw, cleaned and intermediate data are linked "
    "from the relevant tabs and are public in the repository above."
)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(f"wrote {OUT}")
